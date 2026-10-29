"""Compute the open-data figures used in the pitch, straight from the downloaded files.

    python tools/open_data_stats.py

Reads (see data/open/README.md):
  data/open/cause_of_traffic_fines.xlsx                      (Bayanat, Ministry of Interior)
  data/open/dubai_government_service_prices.csv              (Dubai Pulse)
  data/open/unified_federal_traffic_system_2023/Registered vehicles-2023.xlsx  (Bayanat, MOI)

Every number printed here can be re-checked by anyone from the same public files.
"""
import csv
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

OPEN = Path(__file__).resolve().parent.parent / "data" / "open"

# In the Bayanat file this row is a SUBTOTAL of the speed bands above it
# (checked: it equals their sum for every emirate and year), so counting it
# as well would double the speeding numbers.
SPEED_SUBTOTAL = "تجاوز الحد الأقصى للسرعة المقررة"
DUBAI = "دبي"


def traffic_fines(path=OPEN / "cause_of_traffic_fines.xlsx") -> dict:
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    total, speeding = defaultdict(int), defaultdict(int)
    for emirate, year, cause, value in list(ws.iter_rows(values_only=True))[1:]:
        if year is None:
            continue
        key, cause, value = (str(emirate).strip(), int(year)), str(cause).strip(), int(value or 0)
        if cause == SPEED_SUBTOTAL:
            speeding[key] += value
        else:
            total[key] += value
    years = sorted({y for _, y in total})
    out = {}
    for y in years:
        uae_t = sum(v for (e, yy), v in total.items() if yy == y)
        uae_s = sum(v for (e, yy), v in speeding.items() if yy == y)
        out[y] = {"uae_fines": uae_t, "uae_speeding_share": uae_s / uae_t if uae_t else 0,
                  "dubai_fines": total[(DUBAI, y)],
                  "dubai_speeding_share": speeding[(DUBAI, y)] / total[(DUBAI, y)] if total[(DUBAI, y)] else 0}
    return out


def price_list(path=OPEN / "dubai_government_service_prices.csv") -> dict:
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    types = Counter(r.get("pricingtype_en", "").strip() for r in rows)
    fines_by_entity = Counter(r.get("governmententity_en", "").strip() for r in rows
                              if r.get("pricingtype_en", "").strip() == "Fine")
    return {"rows": len(rows), "entities": len({r.get("governmententity_en", "").strip() for r in rows}),
            "by_type": dict(types), "fines_by_entity": dict(fines_by_entity.most_common()),
            "rows_with_amount": sum(1 for r in rows if (r.get("amount_en") or r.get("amount_ar") or "").strip())}


def vehicles_2023(path=OPEN / "unified_federal_traffic_system_2023" / "Registered vehicles-2023.xlsx") -> int:
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["Data Set"]
    return sum(int(r[3]) for r in list(ws.iter_rows(values_only=True))[1:]
               if r and isinstance(r[3], (int, float)))


def main():
    print("== Cause of Traffic Fines (Bayanat, MOI)")
    for y, s in traffic_fines().items():
        print(f"{y}: UAE {s['uae_fines']:,} fines ({s['uae_speeding_share']:.0%} speeding) | "
              f"Dubai {s['dubai_fines']:,} ({s['dubai_speeding_share']:.0%} speeding)")
    p = price_list()
    print("\n== Dubai Government Service Prices (Dubai Pulse)")
    print(f"{p['rows']:,} fee rows from {p['entities']} government entities; by type: {p['by_type']}")
    print(f"Distinct fines: {p['by_type'].get('Fine', 0):,}. Top: {list(p['fines_by_entity'].items())[:5]}")
    print(f"Rows with an amount filled in: {p['rows_with_amount']} (this export has no amounts)")
    print("\n== Registered vehicles 2023 (Bayanat, MOI)")
    print(f"{vehicles_2023():,} vehicles registered in 2023")


if __name__ == "__main__":
    main()
